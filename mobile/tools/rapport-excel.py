#!/usr/bin/env python3
"""RAPPORT DE COMBAT → TABLEAU EXCEL (11/10, demande de l'auteur).

    python3 mobile/tools/rapport-excel.py rapport.json sortie.xlsx

Lit le texte « Données de combat Paradox » (version 1 ou 2) et écrit un classeur lisible :
  Résumé · Combats (1 ligne par combat) · Par étape · Gardiens · Griffes (journal + cumul, version 2).
Les totaux sont des FORMULES (le tableau se recalcule si on corrige une ligne). Heures en heure de Paris.
"""
import json
import sys
from datetime import datetime, timezone
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PARIS = ZoneInfo('Europe/Paris')
RARETE = {'commun': 'C', 'peu_commun': 'PC', 'rare': 'R', 'epique': 'É', 'legendaire': 'L', 'mythique': 'M'}
ISSUE = {'win': 'Victoire', 'lose': 'Défaite', 'quit': 'Abandon'}
MODE = {'aventure': 'Exploration', 'gardien': 'Gardien'}
SOURCE = {'prime': "Prime d'étape", 'defi-jour': 'Défi du jour', 'defi-semaine': 'Défi de la semaine', 'succes': 'Succès',
          'calendrier': 'Calendrier', 'shop-pieces': 'Shop (pièces)', 'shop-diamants': 'Shop (diamants)', 'ascension': 'Ascension',
          'niveau': 'Montée de niveau', 'evolution': 'Évolution', 'dev': 'Dev'}

POLICE = 'Arial'
ENTETE = PatternFill('solid', start_color='2F4F4F')
VERT = PatternFill('solid', start_color='C6EFCE')
ROUGE = PatternFill('solid', start_color='FFC7CE')
GRIS = PatternFill('solid', start_color='F2F2F2')
FIN = Side(style='thin', color='BFBFBF')


def heure(ms):
    if not ms:
        return None
    return datetime.fromtimestamp(ms / 1000, tz=timezone.utc).astimezone(PARIS).replace(tzinfo=None)


def nom(cid):
    return (cid or '?').capitalize()


def equipe_txt(eq):
    return ' · '.join(f"{nom(m.get('id'))} {m.get('niv', '?')}{'*' * (m.get('palier') or 0)} ({RARETE.get(m.get('rarete'), '?')})" for m in (eq or []))


def sorts_ennemis(sorts):
    l = [(k.split('@')[0].replace('ennemi-', ''), v) for k, v in (sorts or {}).items() if k.startswith('ennemi-')]
    return ', '.join(f'{k} ×{v}' for k, v in sorted(l, key=lambda x: -x[1]))


def entete(ws, titres, largeurs):
    for i, (t, w) in enumerate(zip(titres, largeurs), start=1):
        c = ws.cell(row=1, column=i, value=t)
        c.font = Font(name=POLICE, bold=True, color='FFFFFF', size=10)
        c.fill = ENTETE
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = 'A2'


def police_partout(ws):
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if c.font is None or c.font.name != POLICE:
                c.font = Font(name=POLICE, size=10, bold=c.font.bold if c.font else False, color=c.font.color if c.font else None)
            c.border = Border(bottom=FIN)


def convertir(rapport, sortie):
    combats = rapport.get('combats') or []
    gardiens = rapport.get('gardiens') or []
    economie = rapport.get('economie') or []
    wb = Workbook()

    # ---------- Combats ----------
    ws = wb.active
    ws.title = 'Combats'
    titres = ['N°', 'Heure', 'Mode', 'Étape', 'Chap.', 'Boss', 'Essai', 'Résultat', 'Étoiles', 'Ton équipe', 'Niv. moyen équipe',
              'Niv. ennemis', 'Écart (équipe − ennemis)', 'Ennemis', 'Chance annoncée', 'Puissance', 'Parfaits', 'Bien', 'Ratés',
              'Précision', 'Attaques', 'Ripostes', 'Soutiens', 'Sorts ennemis', 'Durée (s)', 'Tours', 'Dégâts infligés',
              'Dégâts reçus', 'PV restants', 'Griffes en poche', 'Poudre', 'Étape (Exploration)']
    largeurs = [5, 15, 11, 7, 6, 6, 6, 10, 7, 42, 9, 9, 11, 40, 9, 9, 7, 6, 6, 9, 8, 8, 8, 30, 8, 6, 10, 10, 9, 9, 7, 9]
    entete(ws, titres, largeurs)
    for n, c in enumerate(combats, start=1):
        r = n + 1
        eq = c.get('equipe') or []
        niv_moy = round(sum((m.get('niv') or 0) for m in eq) / len(eq), 1) if eq else None
        ennemis = c.get('ennemis') or []
        niv_enn = c.get('niveauEnnemis') if c.get('niveauEnnemis') is not None else (ennemis[0].get('niv') if ennemis else None)
        v = c.get('verdicts') or {}
        chance = c.get('chanceAnnoncee')
        mode = MODE.get(c.get('mode'), c.get('mode'))
        ligne = [n, heure(c.get('debut')), mode, c.get('niveau'),
                 c.get('chapitre') if c.get('chapitre') is not None else ((int(c['niveau']) - 1) // 10 + 1 if c.get('niveau') and c.get('mode') == 'aventure' else None),
                 ('Oui' if c.get('boss') else 'Non') if c.get('boss') is not None else None, c.get('essai'),
                 ISSUE.get(c.get('issue'), c.get('issue')), c.get('etoiles'), equipe_txt(eq), niv_moy, niv_enn, None,
                 ' · '.join(f"{nom(e.get('id'))} {e.get('niv', '?')}{'*' * (e.get('palier') or 0)} ({RARETE.get(e.get('rarete'), '?')})" for e in ennemis) or None,
                 (chance / 100) if chance is not None else None, c.get('puissance') or c.get('deckPuissance'),
                 v.get('parfait', 0), v.get('bien', 0), v.get('rate', 0), None, c.get('attaques'), c.get('ripostes'), c.get('soutiens'),
                 sorts_ennemis(c.get('sorts')) or None, c.get('duree'), c.get('tours'), c.get('degatsInfliges'), c.get('degatsRecus'),
                 (c['pvRestantsPct'] / 100) if c.get('pvRestantsPct') is not None else None, c.get('griffes'), c.get('poudre'), None]
        for col, val in enumerate(ligne, start=1):
            ws.cell(row=r, column=col, value=val)
        # Formules (le tableau se recalcule si on corrige une valeur)
        ws.cell(row=r, column=13, value=f'=IF(AND(ISNUMBER(K{r}),ISNUMBER(L{r})),K{r}-L{r},"")')
        ws.cell(row=r, column=20, value=f'=IF(SUM(Q{r}:S{r})>0,Q{r}/SUM(Q{r}:S{r}),"")')
        ws.cell(row=r, column=32, value=f'=IF(C{r}="Exploration",D{r},"")')
        ws.cell(row=r, column=2).number_format = 'dd/mm hh:mm:ss'
        for col in (15, 20, 29):
            ws.cell(row=r, column=col).number_format = '0%'
        ws.cell(row=r, column=13).number_format = '+0.0;-0.0;0'
        ws.cell(row=r, column=10).alignment = Alignment(wrap_text=True, vertical='top')
        ws.cell(row=r, column=14).alignment = Alignment(wrap_text=True, vertical='top')
    der = len(combats) + 1
    if combats:
        ws.conditional_formatting.add(f'H2:H{der}', CellIsRule(operator='equal', formula=['"Victoire"'], fill=VERT))
        ws.conditional_formatting.add(f'H2:H{der}', CellIsRule(operator='equal', formula=['"Défaite"'], fill=ROUGE))
        ws.conditional_formatting.add(f'M2:M{der}', CellIsRule(operator='lessThan', formula=['0'], font=Font(name=POLICE, color='C00000', bold=True)))
        ws.auto_filter.ref = f'A1:{get_column_letter(len(titres))}{der}'
    police_partout(ws)

    # ---------- Par étape (Exploration) ----------
    wp = wb.create_sheet('Par étape')
    entete(wp, ['Étape', 'Essais', 'Victoires', '1er essai', 'Niv. ennemis', 'Niv. équipe (1er essai)', 'Écart (1er essai)',
                'Chance annoncée (1er essai)', 'Étoiles max'], [7, 7, 9, 10, 10, 12, 10, 12, 8])
    etapes = sorted({c.get('niveau') for c in combats if c.get('mode') == 'aventure' and c.get('niveau')})
    rng = f'$2:${der}' if combats else '$2:$2'
    for i, e in enumerate(etapes, start=2):
        wp.cell(row=i, column=1, value=e)
        wp.cell(row=i, column=2, value=f'=COUNTIF(Combats!$AF$2:$AF${der},A{i})')
        wp.cell(row=i, column=3, value=f'=COUNTIFS(Combats!$AF$2:$AF${der},A{i},Combats!$H$2:$H${der},"Victoire")')
        m = f'MATCH(A{i},Combats!$AF$2:$AF${der},0)'
        wp.cell(row=i, column=4, value=f'=IFERROR(INDEX(Combats!$H$2:$H${der},{m}),"")')
        wp.cell(row=i, column=5, value=f'=IFERROR(INDEX(Combats!$L$2:$L${der},{m}),"")')
        wp.cell(row=i, column=6, value=f'=IFERROR(INDEX(Combats!$K$2:$K${der},{m}),"")')
        wp.cell(row=i, column=7, value=f'=IF(AND(ISNUMBER(E{i}),ISNUMBER(F{i})),F{i}-E{i},"")')
        wp.cell(row=i, column=8, value=f'=IFERROR(INDEX(Combats!$O$2:$O${der},{m}),"")')
        wp.cell(row=i, column=9, value=f'=IFERROR(_xlfn.MAXIFS(Combats!$I$2:$I${der},Combats!$AF$2:$AF${der},A{i}),"")')
        wp.cell(row=i, column=8).number_format = '0%'
        wp.cell(row=i, column=7).number_format = '+0.0;-0.0;0'
    fe = len(etapes) + 1
    if etapes:
        wp.conditional_formatting.add(f'D2:D{fe}', CellIsRule(operator='equal', formula=['"Victoire"'], fill=VERT))
        wp.conditional_formatting.add(f'D2:D{fe}', CellIsRule(operator='equal', formula=['"Défaite"'], fill=ROUGE))
    police_partout(wp)
    _ = rng

    # ---------- Gardiens ----------
    wg = wb.create_sheet('Gardiens')
    entete(wg, ['N°', 'Heure', 'Chance annoncée', 'Résultat', 'Chiffre du Gardien', 'Puissance du deck', 'Ton équipe'], [5, 15, 12, 10, 12, 12, 45])
    gc = [c for c in combats if c.get('mode') == 'gardien']
    lignes_g = gc if gc else [{'debut': g.get('t'), 'chanceAnnoncee': g.get('chance'), 'issue': 'win' if g.get('gagne') else 'lose'} for g in gardiens]
    for n, c in enumerate(lignes_g, start=1):
        r = n + 1
        vals = [n, heure(c.get('debut')), (c.get('chanceAnnoncee') / 100) if c.get('chanceAnnoncee') is not None else None,
                ISSUE.get(c.get('issue'), c.get('issue')), c.get('gardienChiffre'), c.get('deckPuissance'), equipe_txt(c.get('equipe'))]
        for col, val in enumerate(vals, start=1):
            wg.cell(row=r, column=col, value=val)
        wg.cell(row=r, column=2).number_format = 'dd/mm hh:mm:ss'
        wg.cell(row=r, column=3).number_format = '0%'
    dg = len(lignes_g) + 1
    if lignes_g:
        wg.conditional_formatting.add(f'D2:D{dg}', CellIsRule(operator='equal', formula=['"Victoire"'], fill=VERT))
        wg.conditional_formatting.add(f'D2:D{dg}', CellIsRule(operator='equal', formula=['"Défaite"'], fill=ROUGE))
    police_partout(wg)

    # ---------- Griffes (journal, version 2) ----------
    wgr = wb.create_sheet('Griffes')
    entete(wgr, ['Heure', 'Source', 'Montant', 'Cumul', 'Étape en cours', 'Détail'], [15, 20, 10, 10, 10, 30])
    chrono = sorted([(c.get('debut') or 0, c.get('niveau')) for c in combats if c.get('mode') == 'aventure'])
    for n, x in enumerate(sorted(economie, key=lambda y: y.get('t') or 0), start=2):
        etape = None
        for t, niv in chrono:
            if t <= (x.get('t') or 0):
                etape = niv
        det = ', '.join(f'{k} {v}' for k, v in x.items() if k not in ('t', 'source', 'montant'))
        wgr.cell(row=n, column=1, value=heure(x.get('t'))).number_format = 'dd/mm hh:mm:ss'
        wgr.cell(row=n, column=2, value=SOURCE.get(x.get('source'), x.get('source')))
        wgr.cell(row=n, column=3, value=x.get('montant'))
        wgr.cell(row=n, column=4, value=f'=C{n}' if n == 2 else f'=D{n - 1}+C{n}')
        wgr.cell(row=n, column=5, value=etape)
        wgr.cell(row=n, column=6, value=det or None)
    fg = len(economie) + 1
    if economie:
        wgr.conditional_formatting.add(f'C2:C{fg}', CellIsRule(operator='lessThan', formula=['0'], font=Font(name=POLICE, color='C00000')))
    else:
        wgr.cell(row=2, column=1, value='Rapport version 1 : pas encore de journal des Griffes (il arrive avec la mise à jour du 11/10).')
    police_partout(wgr)

    # ---------- Résumé (formules) ----------
    wr = wb.create_sheet('Résumé', 0)
    wr.column_dimensions['A'].width = 38
    wr.column_dimensions['B'].width = 14
    wr.column_dimensions['C'].width = 46
    wr['A1'] = 'Rapport de combat Paradox'
    wr['A1'].font = Font(name=POLICE, bold=True, size=14)
    wr['A2'] = f"Exporté le {heure(int(datetime.fromisoformat(rapport['date'].replace('Z', '+00:00')).timestamp() * 1000)).strftime('%d/%m/%Y %H:%M') if rapport.get('date') else '?'} · version du rapport : {rapport.get('version', 1)}"
    wr['A2'].font = Font(name=POLICE, italic=True, size=9, color='595959')
    D = f'{der}'
    lignes = [
        ('EXPLORATION', None, None),
        ('Combats', f'=COUNTIF(Combats!$C$2:$C${D},"Exploration")', ''),
        ('Victoires', f'=COUNTIFS(Combats!$C$2:$C${D},"Exploration",Combats!$H$2:$H${D},"Victoire")', ''),
        ('Taux de victoire', '=IF(B4>0,B5/B4,"")', 'victoires ÷ combats'),
        ('Étapes jouées', f"=COUNT('Par étape'!$A$2:$A${max(fe, 2)})", ''),
        ('Étapes gagnées au 1er essai', f"=COUNTIF('Par étape'!$D$2:$D${max(fe, 2)},\"Victoire\")", ''),
        ('Taux de 1er essai', '=IF(B7>0,B8/B7,"")', 'cible du jeu : ~70 % pour un joueur actif'),
        ('Étape la plus haute', f"=MAX('Par étape'!$A$2:$A${max(fe, 2)})", ''),
        ('Écart moyen équipe − ennemis', f'=IFERROR(AVERAGE(Combats!$M$2:$M${D}),"")', 'en niveaux (négatif = ennemis au-dessus)'),
        ('Durée moyenne d\'un combat (s)', f'=IFERROR(AVERAGEIF(Combats!$C$2:$C${D},"Exploration",Combats!$Y$2:$Y${D}),"")', ''),
        ('PRÉCISION À LA JAUGE', None, None),
        ('Parfaits', f'=SUM(Combats!$Q$2:$Q${D})', ''),
        ('Bien', f'=SUM(Combats!$R$2:$R${D})', ''),
        ('Ratés', f'=SUM(Combats!$S$2:$S${D})', ''),
        ('Part de parfaits', '=IF(SUM(B14:B16)>0,B14/SUM(B14:B16),"")', 'joueur de référence du jeu : 47 %'),
        ('GARDIENS', None, None),
        ('Combats de Gardien', f'=COUNTA(Gardiens!$A$2:$A${max(dg, 2)})', ''),
        ('Gagnés', f'=COUNTIF(Gardiens!$D$2:$D${max(dg, 2)},"Victoire")', ''),
        ('Chance annoncée moyenne', f'=IFERROR(AVERAGE(Gardiens!$C$2:$C${max(dg, 2)}),"")', 'le jeu vise ~80 %'),
        ('GRIFFES (journal)', None, None),
    ]
    r = 3
    for lib, f, note in lignes:
        wr.cell(row=r, column=1, value=lib)
        if f is None:
            wr.cell(row=r, column=1).font = Font(name=POLICE, bold=True, color='FFFFFF')
            for col in (1, 2, 3):
                wr.cell(row=r, column=col).fill = ENTETE
        else:
            wr.cell(row=r, column=1).font = Font(name=POLICE)
            wr.cell(row=r, column=2, value=f).font = Font(name=POLICE, bold=True)
            wr.cell(row=r, column=3, value=note or None).font = Font(name=POLICE, italic=True, size=9, color='595959')
        r += 1
    for rr in (6, 9, 17, 21):
        wr.cell(row=rr, column=2).number_format = '0%'
    wr.cell(row=11, column=2).number_format = '+0.0;-0.0;0'
    wr.cell(row=12, column=2).number_format = '0'
    for src in ['prime', 'defi-jour', 'defi-semaine', 'succes', 'calendrier', 'shop-pieces', 'shop-diamants', 'ascension', 'niveau', 'evolution']:
        wr.cell(row=r, column=1, value=SOURCE[src]).font = Font(name=POLICE)
        wr.cell(row=r, column=2, value=f'=SUMIF(Griffes!$B$2:$B${max(fg, 2)},A{r},Griffes!$C$2:$C${max(fg, 2)})').font = Font(name=POLICE, bold=True)
        r += 1
    wr.cell(row=r, column=1, value='Solde (gains − dépenses)').font = Font(name=POLICE, bold=True)
    wr.cell(row=r, column=2, value=f'=SUM(B{r - 10}:B{r - 1})').font = Font(name=POLICE, bold=True)
    wr.cell(row=r + 2, column=1, value="Les totaux sont des formules : corriger une ligne dans « Combats » met tout à jour.").font = Font(name=POLICE, italic=True, size=9, color='595959')
    wb.save(sortie)


if __name__ == '__main__':
    with open(sys.argv[1], encoding='utf-8') as f:
        convertir(json.load(f), sys.argv[2])
    print('classeur écrit :', sys.argv[2])
